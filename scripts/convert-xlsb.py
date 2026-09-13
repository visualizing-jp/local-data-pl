#!/usr/bin/env python3
"""資料集が xlsb だけのとき、値だけ xlsx に直す。pyxlsb と openpyxl が要る。"""

import sys

from openpyxl import Workbook
from pyxlsb import open_workbook


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: convert-xlsb.py SRC.xlsb DEST.xlsx")
    src, dest = sys.argv[1], sys.argv[2]
    out = Workbook()
    out.remove(out.active)
    with open_workbook(src) as wb:
        for name in wb.sheets:
            ws_out = out.create_sheet(title=str(name)[:31])
            with wb.get_sheet(name) as ws:
                for row in ws.rows():
                    for cell in row:
                        if cell.v is None:
                            continue
                        ws_out.cell(cell.r + 1, cell.c + 1, cell.v)
    out.save(dest)


if __name__ == "__main__":
    main()
