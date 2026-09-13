#!/usr/bin/env python3
"""資料集が xls だけのとき、値だけ xlsx に直す。xlrd と openpyxl が要る。"""

import sys

import xlrd
from openpyxl import Workbook


def cell_value(cell, datemode: int):
    if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    return cell.value


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: convert-xls.py SRC.xls DEST.xlsx")
    src, dest = sys.argv[1], sys.argv[2]
    book = xlrd.open_workbook(src, formatting_info=False)
    out = Workbook()
    out.remove(out.active)
    for sheet in book.sheets():
        ws_out = out.create_sheet(title=str(sheet.name)[:31])
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                value = cell_value(sheet.cell(r, c), book.datemode)
                if value is None or value == "":
                    continue
                ws_out.cell(r + 1, c + 1, value)
    out.save(dest)


if __name__ == "__main__":
    main()
